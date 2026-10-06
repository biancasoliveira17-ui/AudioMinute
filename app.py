import io
import os
import re

import streamlit as st
from openai import OpenAI
from docx import Document


st.set_page_config(
    page_title="AudioMinute",
    page_icon="🎙️",
    layout="wide"
)

st.markdown(
    """
    <style>
    .main-title { font-size: 42px; font-weight: 700; margin-bottom: 0; }
    .subtitle { font-size: 18px; color: #666; margin-top: 5px; margin-bottom: 25px; }
    </style>
    """,
    unsafe_allow_html=True
)

st.markdown('<div class="main-title">🎙️ AudioMinute</div>', unsafe_allow_html=True)
st.markdown(
    '<div class="subtitle">Transforme reuniões em atas e planos de ação com Inteligência Artificial</div>',
    unsafe_allow_html=True
)

st.write(
    "O AudioMinute transforma gravações de reuniões em informações organizadas, "
    "identificando tópicos, decisões, responsáveis, tarefas e prazos."
)
st.divider()


def obter_api_key():
    try:
        if "OPENAI_API_KEY" in st.secrets:
            return st.secrets["OPENAI_API_KEY"]
    except Exception:
        pass
    return os.getenv("OPENAI_API_KEY")


api_key = obter_api_key()

if not api_key:
    st.error(
        """⚠️ A chave da OpenAI não foi configurada.

No Streamlit Cloud, configure em Settings → Secrets:

OPENAI_API_KEY = "sua-chave"

Para executar localmente, configure a variável de ambiente OPENAI_API_KEY."""
    )
    st.stop()

client = OpenAI(api_key=api_key)


def transcrever_audio(arquivo):
    arquivo.seek(0)
    resultado = client.audio.transcriptions.create(
        model="whisper-1",
        file=(arquivo.name, arquivo.getvalue(), arquivo.type)
    )
    return resultado.text


def gerar_ata(transcricao):
    prompt = f"""
Você é um assistente especializado em elaboração de atas de reuniões corporativas.

Analise cuidadosamente a transcrição abaixo.

Crie uma ata profissional e objetiva.

Utilize EXATAMENTE esta estrutura:

# TÍTULO DA REUNIÃO

# DATA

# RESUMO

# PRINCIPAIS TÓPICOS

- tópico 1
- tópico 2
- tópico 3

# DECISÕES TOMADAS

- decisão 1
- decisão 2

# PLANO DE AÇÃO

| Responsável | Tarefa | Prazo |
|-------------|--------|-------|
| Nome | Tarefa | Prazo |

REGRAS IMPORTANTES:
1. Não invente informações.
2. Utilize somente informações presentes na transcrição.
3. Caso uma informação não esteja disponível, escreva "Não informado."
4. Identifique corretamente responsáveis, tarefas e prazos.
5. Seja objetivo e profissional.
6. Não crie participantes que não aparecem na transcrição.

TRANSCRIÇÃO DA REUNIÃO:

{transcricao}
"""

    resposta = client.chat.completions.create(
        model="gpt-4o-mini",
        messages=[
            {
                "role": "system",
                "content": "Você é um especialista em atas e produtividade corporativa."
            },
            {"role": "user", "content": prompt}
        ],
        temperature=0.2
    )
    return resposta.choices[0].message.content


def criar_docx(texto):
    documento = Document()
    documento.add_heading("AudioMinute - Ata da Reunião", level=1)

    linhas = texto.splitlines()
    tabela_linhas = []

    for linha in linhas:
        linha = linha.strip()
        if not linha:
            continue

        if linha.startswith("# "):
            documento.add_heading(linha.replace("# ", "").strip(), level=2)
        elif linha.startswith("|"):
            if "---" not in linha:
                valores = [item.strip() for item in linha.strip("|").split("|")]
                if valores:
                    tabela_linhas.append(valores)
        elif linha.startswith("- "):
            documento.add_paragraph(linha[2:].strip(), style="List Bullet")
        else:
            documento.add_paragraph(linha)

    if tabela_linhas:
        tabela = documento.add_table(rows=1, cols=len(tabela_linhas[0]))
        tabela.style = "Table Grid"

        for i, valor in enumerate(tabela_linhas[0]):
            tabela.rows[0].cells[i].text = valor

        for linha in tabela_linhas[1:]:
            cells = tabela.add_row().cells
            for i, valor in enumerate(linha):
                if i < len(cells):
                    cells[i].text = valor

    arquivo = io.BytesIO()
    documento.save(arquivo)
    arquivo.seek(0)
    return arquivo


def nome_seguro(nome):
    return re.sub(r"[^a-zA-Z0-9_.-]", "_", nome)


st.header("1. 🎧 Envie o áudio da reunião")
st.markdown(
    "**Formatos aceitos:** MP3, WAV, M4A e MPEG\n\n"
    "Para a demonstração, recomendamos uma gravação de aproximadamente 1 a 2 minutos."
)

arquivo_audio = st.file_uploader(
    "Selecione o arquivo de áudio",
    type=["mp3", "wav", "m4a", "mpeg"]
)

if arquivo_audio:
    st.audio(arquivo_audio, format=arquivo_audio.type)
    st.success(f"✅ Arquivo carregado: {arquivo_audio.name}")

    tamanho_mb = arquivo_audio.size / (1024 * 1024)
    st.caption(f"Tamanho do arquivo: {tamanho_mb:.2f} MB")

    st.divider()

    if st.button("🚀 Processar reunião", type="primary", use_container_width=True):
        try:
            st.subheader("🎧 Etapa 1 — Transcrição")

            with st.spinner("Transcrevendo o áudio..."):
                transcricao = transcrever_audio(arquivo_audio)

            st.success("✅ Áudio transcrito com sucesso!")

            with st.expander("📝 Visualizar transcrição"):
                st.write(transcricao)

            st.subheader("🤖 Etapa 2 — Inteligência Artificial")

            with st.spinner("A Inteligência Artificial está organizando a reunião..."):
                ata = gerar_ata(transcricao)

            st.success("✅ Ata criada com sucesso!")

            st.divider()
            st.header("2. 📋 Ata da reunião")
            st.markdown(ata)

            st.divider()
            st.header("3. 📥 Exportar resultado")

            col1, col2 = st.columns(2)

            with col1:
                st.download_button(
                    label="📥 Baixar ata em TXT",
                    data=ata.encode("utf-8"),
                    file_name="AudioMinute_Ata.txt",
                    mime="text/plain",
                    use_container_width=True
                )

            with col2:
                arquivo_docx = criar_docx(ata)
                st.download_button(
                    label="📄 Baixar ata em DOCX",
                    data=arquivo_docx,
                    file_name="AudioMinute_Ata.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    use_container_width=True
                )

        except Exception as erro:
            st.error("❌ Ocorreu um erro durante o processamento.")
            with st.expander("Detalhes técnicos"):
                st.exception(erro)

st.divider()
st.caption(
    "AudioMinute — Gerador de Atas e Tarefas a partir de Reuniões | "
    "Projeto de Inteligência Artificial | Líder: Ayila Melissa"
)
